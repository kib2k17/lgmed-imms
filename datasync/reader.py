"""
Reading a monitoring workbook into rows a sync can act on.

The office's registers are spreadsheets kept by hand over a year, and they
drift: a column is renamed in April, a header cell is left empty from June, a
day's date is written once and the rows beneath it are left blank. The reader
absorbs that drift so the engine never sees it. It knows nothing about any
model - it turns sheets into `ParsedRow`s and `SkippedRow`s according to a
`SyncProfile`, and says, per sheet, which column it took each field from.

Nothing here writes to the database.
"""

import datetime
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

# Written dates the registers have been seen to use, most common first.
DATE_FORMATS = (
    "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y",
    "%B %d, %Y", "%B %d %Y", "%b %d, %Y", "%b %d %Y",
)

# Excel stores dates as days since 1899-12-30. A bare number in a date column
# inside this range (1954 to 2119) is taken to be one; outside it, it is not.
EXCEL_SERIAL_RANGE = (20000, 80000)


def normalise_header(value):
    """`Assigned Focal / Remarks` and `ASSIGNED FOCAL/REMARKS` are the same header."""
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def clean_text(value, multiline=False):
    """A cell as text: whitespace tidied, numbers without a stray `.0`."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime.datetime):
        value = value.date()
    if isinstance(value, datetime.date):
        return value.isoformat()
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    if multiline:
        lines = (" ".join(line.split()) for line in text.split("\n"))
        return "\n".join(line for line in lines if line)
    return " ".join(text.split())


def parse_date(value):
    """
    A cell as a date, or None when it cannot be read as one.

    Returns (date, ok). `ok` is False only for a cell that holds something but
    not a date - an empty cell is (None, True), so the caller can tell "no date
    written here" from "a date written here that makes no sense".
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if isinstance(value, datetime.datetime):
        return value.date(), True
    if isinstance(value, datetime.date):
        return value, True
    if isinstance(value, (int, float)):
        low, high = EXCEL_SERIAL_RANGE
        if low <= value <= high:
            return from_excel(value).date(), True
        return None, False
    # "Feb.12, 2026" and "Sept. 3, 2026" as written, read as "Feb 12, 2026".
    text = re.sub(r"\.\s*", " ", str(value))
    text = re.sub(r"\bsept\b", "Sep", " ".join(text.split()), flags=re.I)
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date(), True
        except ValueError:
            continue
    return None, False


@dataclass
class SheetReport:
    """What happened to one sheet, for the preview and the summary."""

    name: str
    status: str = "read"          # read | skipped | excluded
    reason: str = ""
    header_row: int | None = None
    columns: list = field(default_factory=list)   # [(field label, "B: DOCKET NUMBER")]
    # The same, spelt out for a migration's column check: one dict per field,
    # saying how its column was found (see `_describe_mapping`).
    mapping: list = field(default_factory=list)
    rows: int = 0                 # rows that became records to sync
    skipped: int = 0              # rows with content that could not be used
    blank: int = 0                # empty rows passed over

    def as_dict(self):
        return {
            "name": self.name, "status": self.status, "reason": self.reason,
            "header_row": self.header_row, "columns": self.columns,
            "mapping": self.mapping,
            "rows": self.rows, "skipped": self.skipped, "blank": self.blank,
        }


@dataclass
class ParsedRow:
    sheet: str
    row: int
    key: str
    date: datetime.date | None
    values: dict                  # column name -> cleaned text
    notes: list = field(default_factory=list)   # kept, but worth a second look
    # column name -> "date" | "number" for a cell Excel stored as something
    # other than text, so a target can notice a reference Excel reformatted.
    raw_kinds: dict = field(default_factory=dict)
    # What a target worked out while checking the row, for the preview.
    extra: dict = field(default_factory=dict)

    @property
    def location(self):
        return f"{self.sheet} row {self.row}"


@dataclass
class SkippedRow:
    sheet: str
    row: int
    key: str
    reason: str
    # error: could not be used | duplicate: already read or already stored |
    # ignored: not a record at all (the template's sample row)
    kind: str = "error"
    values: dict = field(default_factory=dict)

    def as_dict(self):
        return {"sheet": self.sheet, "row": self.row, "key": self.key,
                "reason": self.reason, "kind": self.kind}


@dataclass
class ReadResult:
    sheets: list
    rows: list
    skipped: list

    @property
    def flagged(self):
        return [row for row in self.rows if row.notes]

    @property
    def sheet_names(self):
        return [sheet.name for sheet in self.sheets]


# ---------------------------------------------------------------------------
# Headers
# ---------------------------------------------------------------------------


def _match_headers(profile, header):
    """
    Which column each field comes from, judged by header text alone.

    An exact match is preferred to a prefix match, so `DATE` is never taken by
    a `DATE SENT` column when a `DATE` column is also present.
    """
    normalised = [normalise_header(cell) for cell in header]
    claimed = {}
    for exact in (True, False):
        for column in profile.columns:
            if column.name in claimed:
                continue
            for alias in column.normalised_aliases:
                for index, text in enumerate(normalised):
                    if not text or index in claimed.values():
                        continue
                    if text == alias if exact else text.startswith(alias):
                        claimed[column.name] = index
                        break
                if column.name in claimed:
                    break
    return claimed, normalised


def _resolve_columns(profile, header):
    """
    The column each field is read from, including the ones found by position.

    A field marked `blank_header_fallback` whose header cell is empty is taken
    from the leftmost column with an empty header - to the left of the key
    column, unless it is the key itself. This is how the registers' Date column
    is found in the months its header cell was left blank.
    """
    claimed, normalised = _match_headers(profile, header)
    by_position = set()

    fallbacks = sorted(
        (c for c in profile.columns if c.blank_header_fallback),
        key=lambda c: c.name != profile.key_column,   # the key first
    )
    for column in fallbacks:
        if column.name in claimed:
            continue
        key_index = claimed.get(profile.key_column)
        for index, text in enumerate(normalised):
            if text or index in claimed.values():
                continue
            if column.name != profile.key_column and (
                key_index is None or index > key_index
            ):
                continue
            claimed[column.name] = index
            by_position.add(column.name)
            break

    # Unlabelled columns straight after a field that collects them - the
    # "Rerouted to PDMU" written beside a remark.
    extras = {}
    taken = set(claimed.values())
    for column in profile.columns:
        if not column.append_unlabelled or column.name not in claimed:
            continue
        following = []
        index = claimed[column.name] + 1
        while index < len(normalised) and not normalised[index] and index not in taken:
            following.append(index)
            index += 1
        extras[column.name] = following
    return claimed, by_position, extras


def _find_header_row(profile, rows):
    """The first row, within the profile's reach, that names enough fields."""
    for index, row in enumerate(rows[: profile.header_scan_rows]):
        claimed, _ = _match_headers(profile, row)
        if len(claimed) >= profile.min_header_matches:
            return index
    return None


def _describe_columns(profile, header, claimed, by_position, extras):
    described = []
    for column in profile.columns:
        if column.name not in claimed:
            described.append((column.label, "not found"))
            continue
        index = claimed[column.name]
        text = clean_text(header[index]) if index < len(header) else ""
        where = f"{get_column_letter(index + 1)}: {text}" if text else (
            f"{get_column_letter(index + 1)} (blank header)"
        )
        if column.name in by_position:
            where += ", by position"
        if extras.get(column.name):
            letters = ", ".join(get_column_letter(i + 1) for i in extras[column.name])
            where += f" + unlabelled {letters}"
        described.append((column.label, where))
    return described


def _describe_mapping(profile, header, claimed, by_position):
    """
    Each field's column, and how it was recognised.

    `how` is "exact" when the header is the expected one give or take case,
    spacing and punctuation; "alias" when it is one of the other names the
    office has used for that column, which the person should see and confirm;
    "position" for a blank header found by where it sits; "missing" otherwise.
    """
    mapping = []
    for column in profile.columns:
        entry = {
            "field": column.name, "label": column.label,
            "expected": column.aliases[0],
            "required": column.name in profile.required_columns,
            "header": "", "column": "", "how": "missing",
        }
        if column.name in claimed:
            index = claimed[column.name]
            text = clean_text(header[index]) if index < len(header) else ""
            entry["header"] = text
            entry["column"] = get_column_letter(index + 1)
            if column.name in by_position:
                entry["how"] = "position"
            elif normalise_header(text) == column.normalised_aliases[0]:
                entry["how"] = "exact"
            else:
                entry["how"] = "alias"
        mapping.append(entry)
    return mapping


def _raw_kind(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    return None


# ---------------------------------------------------------------------------
# Rows
# ---------------------------------------------------------------------------


def _cell(row, index):
    return row[index] if index is not None and index < len(row) else None


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def read_workbook(path_or_file, profile, sheets=None):
    """
    Read every sheet of a workbook against a profile.

    `sheets`, when given, is the set of sheet names to read; the others are
    reported as excluded rather than silently dropped, so the summary still
    accounts for the whole file.
    """
    workbook = load_workbook(path_or_file, read_only=True, data_only=True)
    reports, parsed, skipped = [], [], []
    seen = {}     # folded key -> the ParsedRow it first appeared as
    try:
        for worksheet in workbook.worksheets:
            report = SheetReport(name=worksheet.title)
            reports.append(report)
            if sheets is not None and worksheet.title not in sheets:
                report.status, report.reason = "excluded", "Not selected for this sync"
                continue
            rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
            _read_sheet(profile, report, rows, parsed, skipped, seen)
    finally:
        workbook.close()
    return ReadResult(sheets=reports, rows=parsed, skipped=skipped)


def _read_sheet(profile, report, rows, parsed, skipped, seen):
    if not any(not _is_blank(value) for row in rows for value in row):
        report.status, report.reason = "skipped", "Empty sheet"
        return

    header_index = _find_header_row(profile, rows)
    if header_index is None:
        report.status = "skipped"
        report.reason = (
            f"No header row found in the first {profile.header_scan_rows} "
            f"row{'s' if profile.header_scan_rows > 1 else ''}"
        )
        return

    header = rows[header_index]
    claimed, by_position, extras = _resolve_columns(profile, header)
    report.header_row = header_index + 1
    report.columns = _describe_columns(profile, header, claimed, by_position, extras)
    report.mapping = _describe_mapping(profile, header, claimed, by_position)

    key_label = profile.column(profile.key_column).label
    if profile.key_column not in claimed:
        report.status, report.reason = "skipped", f"No {key_label} column"
        return
    missing = [profile.column(name).label for name in profile.required_columns
               if name not in claimed]
    if missing:
        report.status = "skipped"
        report.reason = f"Required column not found: {', '.join(missing)}"
        return

    ignore = re.compile(profile.ignore_key_pattern, re.I) if profile.ignore_key_pattern else None
    from_key = re.compile(profile.date_from_key) if profile.date_from_key else None
    date_index = claimed.get(profile.date_column) if profile.date_column else None
    fold = profile.get_key_normaliser()
    last_date = None

    def skip(row_number, key, reason, kind="error", values=None):
        report.skipped += 1
        skipped.append(
            SkippedRow(report.name, row_number, key, reason, kind, values or {})
        )

    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        values = {}
        raw_kinds = {}
        for column in profile.columns:
            if column.name == profile.date_column or column.name not in claimed:
                continue
            raw = _cell(row, claimed[column.name])
            if _raw_kind(raw):
                raw_kinds[column.name] = _raw_kind(raw)
            text = clean_text(raw, column.multiline)
            for index in extras.get(column.name, ()):
                extra = clean_text(_cell(row, index), column.multiline)
                if extra:
                    text = f"{text}\n{extra}" if text else extra
            values[column.name] = text

        raw_date = _cell(row, date_index)
        if all(not text for text in values.values()):
            # Nothing but (perhaps) a date: a blank row, or a day heading with
            # no entries under it. It still moves the carried date along.
            date, ok = parse_date(raw_date)
            if date and ok:
                last_date = date
            report.blank += 1
            continue

        key = values.get(profile.key_column, "")
        if not key:
            skip(offset, "", f"No {key_label}", values=values)
            continue
        if ignore and ignore.search(key):
            skip(offset, key, "Sample / template row", kind="ignored", values=values)
            continue

        date = None
        notes = []
        if date_index is not None:
            date, ok = parse_date(raw_date)
            if not ok:
                # A name or a note written in the date column. The entry is
                # real; the date is the day it sits under, and the text is
                # kept with the remarks rather than lost.
                stray = clean_text(raw_date)
                date = last_date if profile.carry_date_forward else None
                if profile.remarks_column:
                    kept = f"; the text was kept in {profile.column(profile.remarks_column).label}"
                else:
                    kept = ""
                used = f"used {date:%d %b %Y} from the row above" if date else "no date to use"
                notes.append(f"Date column held {stray!r}: {used}{kept}")
                if profile.remarks_column:
                    remark = values.get(profile.remarks_column, "")
                    values[profile.remarks_column] = (
                        f"{remark}\n{stray}" if remark else stray
                    )
            elif date:
                last_date = date
            elif profile.carry_date_forward:
                date = last_date
        if date is None and from_key:
            match = from_key.search(key)
            if match:
                try:
                    date = datetime.date(*(int(part) for part in match.groups()))
                except ValueError:
                    date = None
        if date is None and profile.date_required:
            skip(offset, key, "No date on this row or any row above it on the sheet",
                 values=values)
            continue

        folded = fold(key)
        if folded in seen:
            first = seen[folded]
            written = "" if first.key == key else f" (written there as {first.key})"
            if profile.merge_duplicates:
                merged = _merge_duplicate(profile, first, values)
                skip(
                    offset, key,
                    f"Duplicate {key_label} - already read from {first.location}{written}"
                    + (f"; merged its {', '.join(merged)} into that row" if merged else ""),
                    kind="duplicate", values=values,
                )
            elif _same_entry(profile, first.values, values):
                skip(offset, key,
                     f"Duplicate {key_label} - the same entry as {first.location}{written}",
                     kind="duplicate", values=values)
            else:
                # One code on two different communications. Only one record
                # can hold a code, so this is for the person to resolve.
                skip(offset, key,
                     f"{key_label}: also used on {first.location}{written} for a "
                     "different communication - give this row its own code",
                     values=values)
            continue

        row = ParsedRow(report.name, offset, key, date, values, notes, raw_kinds)
        seen[folded] = row
        parsed.append(row)
        report.rows += 1


def _same_entry(profile, first, values):
    """Whether a repeated row says nothing the first did not, ignoring spelling."""
    return all(
        normalise_header(values.get(column.name, ""))
        in normalise_header(first.get(column.name, ""))
        for column in profile.columns
    )


def _merge_duplicate(profile, first, values):
    """
    Fold what a repeated entry adds into the first one.

    The registers repeat a docket when a document comes back or is re-routed,
    often with a newer instruction. The first entry keeps its date - that is
    when the document arrived - but a remark the repeat adds is appended and a
    field the first entry left blank is filled, so nothing written is lost.
    """
    merged = []
    for column in profile.columns:
        new = values.get(column.name, "")
        old = first.values.get(column.name, "")
        if not new or column.name == profile.key_column:
            continue
        # "NAOMI - PLS ACT ON THIS" repeated as "Naomi- Pls. act on this" adds
        # nothing, so compare without case, spacing or punctuation.
        if normalise_header(new) in normalise_header(old):
            continue
        if not old:
            first.values[column.name] = new
        elif column.multiline:
            first.values[column.name] = f"{old}\n{new}"
        else:
            continue
        merged.append(column.label)
    return merged
