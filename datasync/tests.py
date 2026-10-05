"""
Tests for Data Sync.

The workbooks here are built to reproduce the drift in the office's real
registers - a blank Date header, an empty second row, a renamed column, a date
written only on the first row of the day, a remark spilling into an unlabelled
column - because those are what the sync exists to absorb.
"""

import datetime
import io
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from accounts.models import Role, User
from audit.models import Action, AuditEvent
from incoming.models import IncomingDocument, IncomingEvent, IncomingStatus
from outgoing.models import OutgoingDocument

from . import engine
from .models import BatchStatus, SyncBatch
from .profiles import get_profile
from .reader import read_workbook

D = datetime.datetime


def incoming_workbook():
    book = Workbook()
    january = book.active
    january.title = "JANUARY"
    january.append([" ", "HUC/Province", "DOCKET NUMBER", "SUBJECT",
                    "ASSIGNED FOCAL/RMEARKS", None])
    january.append([D(2026, 1, 5), "ADN", "2026-01-05-001", "First subject",
                    "Karen- Pls act", "Rerouted to PDMU"])
    january.append([None, None, "2026-01-05-002", "Second subject", "Dave- Pls act"])
    january.append([None, None, None, None, None])
    january.append([D(2026, 1, 6), None, None, "A subject without a docket", "x"])
    january.append([None, None, "2026-01-05-001", "First subject", "Karen- later note"])

    february = book.create_sheet("FEBRUARY")
    february.append(["Date", "HUC/Province", "DOCKET NUMBER", "SUBJECT",
                     "ASSIGNED FOCAL/RMEARKS"])
    february.append([None, None, None, None, None])       # the empty row 2
    february.append(["Feb.12, 2026", None, "2026-02-12-001", "Feb subject", "Jen"])

    book.create_sheet("Sheet1")

    april = book.create_sheet("APRIL")
    april.append(["DATE", "DOCKET NUMBER", "DESCRIPTION", "CONCERN PERSON"])
    april.append([D(2026, 4, 1), "2026-04-01-001", "April description", "Fritz"])
    april.append(["lynn", "2026-04-01-002", "Name in the date column", "Lynn- act"])

    june = book.create_sheet("JUNE2026")
    june.append(["  ", "DOCKET NUMBER", "DESCRIPTION", "CONCERN PERSON"])
    june.append([None, None, None, None])
    june.append([D(2026, 6, 1), "2026-06-01-001", "June description", "Naomi"])
    june.append([None, "2026-06-01-002", "Carried date", None])
    return book


def as_upload(book, name="2026 Incoming Monitoring.xlsx"):
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleUploadedFile(
        name, buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def as_stream(book):
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


class ReaderTests(TestCase):
    def setUp(self):
        self.result = read_workbook(as_stream(incoming_workbook()), get_profile("incoming"))
        self.rows = {row.key: row for row in self.result.rows}

    def test_every_monthly_sheet_is_read_and_empty_sheets_skipped(self):
        status = {sheet.name: sheet.status for sheet in self.result.sheets}
        self.assertEqual(status["Sheet1"], "skipped")
        for name in ("JANUARY", "FEBRUARY", "APRIL", "JUNE2026"):
            self.assertEqual(status[name], "read", name)

    def test_blank_date_header_is_found_by_position(self):
        self.assertEqual(self.rows["2026-01-05-001"].date, datetime.date(2026, 1, 5))
        self.assertEqual(self.rows["2026-06-01-001"].date, datetime.date(2026, 6, 1))

    def test_date_is_carried_forward_to_rows_below(self):
        self.assertEqual(self.rows["2026-01-05-002"].date, datetime.date(2026, 1, 5))
        self.assertEqual(self.rows["2026-06-01-002"].date, datetime.date(2026, 6, 1))

    def test_written_dates_are_read(self):
        self.assertEqual(self.rows["2026-02-12-001"].date, datetime.date(2026, 2, 12))

    def test_april_layout_maps_description_and_concern_person(self):
        row = self.rows["2026-04-01-001"]
        self.assertEqual(row.values["subject"], "April description")
        self.assertEqual(row.values["remarks"], "Fritz")
        self.assertNotIn("source", row.values)

    def test_unlabelled_remark_is_appended(self):
        remarks = self.rows["2026-01-05-001"].values["remarks"]
        self.assertIn("Karen- Pls act", remarks)
        self.assertIn("Rerouted to PDMU", remarks)

    def test_row_without_docket_is_skipped_with_reason(self):
        reasons = [row.reason for row in self.result.skipped]
        self.assertIn("No Docket Number", reasons)

    def test_repeated_docket_is_skipped_but_its_remark_kept(self):
        duplicate = [row for row in self.result.skipped if row.key == "2026-01-05-001"]
        self.assertEqual(len(duplicate), 1)
        self.assertIn("Duplicate", duplicate[0].reason)
        self.assertIn("Karen- later note", self.rows["2026-01-05-001"].values["remarks"])

    def test_text_in_the_date_column_is_kept_and_flagged(self):
        row = self.rows["2026-04-01-002"]
        self.assertEqual(row.date, datetime.date(2026, 4, 1))
        self.assertIn("lynn", row.values["remarks"])
        self.assertTrue(row.notes)

    def test_blank_rows_are_not_reported_as_skipped(self):
        self.assertEqual(len(self.result.skipped), 2)   # no docket, and the repeat
        self.assertGreater(sum(sheet.blank for sheet in self.result.sheets), 0)


class OutgoingReaderTests(TestCase):
    def test_header_below_a_title_block_and_date_from_control_code(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "2025"
        sheet.append([" ", "   "])
        sheet.append(["LGMED 13"])
        sheet.append(["Control Code", "Type of Communication / Documents",
                      "Subject / Title", "Sent to", "Sent Via", "Remarks",
                      "DMS Number\n(Incoming)", "DMS Number \n(Outgoing)"])
        sheet.append(["Sample:\nLGMED-13 - (RGFJ) - 2023-01-03-0001", "Memo",
                      "Road Clearing", "PDs/CD", "Email", None, "DMS No: ", "DMS No: "])
        sheet.append(["LGMED-13 - (DNB) - 2025-01-02-0001", "Letter", "A subject",
                      "NBOO", "DMS", None, "2024-12-20-001", None])
        result = read_workbook(as_stream(book), get_profile("outgoing"))

        self.assertEqual(result.sheets[0].header_row, 3)
        self.assertEqual(len(result.rows), 1)
        row = result.rows[0]
        self.assertEqual(row.date, datetime.date(2025, 1, 2))
        self.assertEqual(row.values["incoming_reference"], "2024-12-20-001")
        self.assertEqual(result.skipped[0].reason, "Sample / template row")


class SyncTestCase(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._protected = tempfile.mkdtemp()
        cls._override = override_settings(PROTECTED_MEDIA_ROOT=cls._protected)
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._protected, ignore_errors=True)
        super().tearDownClass()

    @classmethod
    def setUpTestData(cls):
        cls.chief = User.objects.create_user(
            username="chief", password="pw", role=Role.ADMIN,
            first_name="Divina", last_name="Cruz",
        )
        cls.encoder = User.objects.create_user(
            username="encoder", password="pw", role=Role.ENCODER,
            first_name="Elena", last_name="Reyes",
        )

    def upload(self, book, module="incoming", user=None):
        self.client.force_login(user or self.chief)
        return self.client.post(
            reverse("datasync:home"), {"module": module, "file": as_upload(book)}
        )


class EngineTests(SyncTestCase):
    def make_batch(self, book):
        self.upload(book)
        return SyncBatch.objects.latest("pk")

    def test_preview_writes_nothing(self):
        batch = self.make_batch(incoming_workbook())
        self.assertEqual(batch.status, BatchStatus.PREVIEW)
        self.assertEqual(IncomingDocument.objects.count(), 0)
        self.assertEqual(batch.counts["created"], 7)

    def test_commit_creates_imported_records_with_a_trail(self):
        batch = self.make_batch(incoming_workbook())
        engine.commit(batch, self.chief)

        self.assertEqual(IncomingDocument.objects.count(), 7)
        document = IncomingDocument.objects.get(docket_number="2026-01-05-002")
        self.assertEqual(document.status, IncomingStatus.IMPORTED)
        self.assertEqual(document.date_received, datetime.date(2026, 1, 5))
        self.assertEqual(document.document_type.name, "Unclassified")
        self.assertEqual(document.created_by, self.chief)
        event = IncomingEvent.objects.get(document=document)
        self.assertIn("JANUARY row 3", event.notes)
        self.assertTrue(AuditEvent.objects.filter(action=Action.SYNC).exists())

    def test_resync_updates_rather_than_duplicates(self):
        engine.commit(self.make_batch(incoming_workbook()), self.chief)

        book = incoming_workbook()
        book["JUNE2026"]["C3"] = "June description, corrected"
        book["JUNE2026"]["D4"] = None     # a blank cell must not erase anything
        plan = engine.commit(self.make_batch(book), self.chief)

        self.assertEqual(IncomingDocument.objects.count(), 7)
        self.assertEqual(plan.counts["created"], 0)
        self.assertEqual(plan.counts["updated"], 1)
        self.assertEqual(
            IncomingDocument.objects.get(docket_number="2026-06-01-001").subject,
            "June description, corrected",
        )

    def test_blank_cell_does_not_erase_a_stored_value(self):
        engine.commit(self.make_batch(incoming_workbook()), self.chief)
        book = incoming_workbook()
        book["JUNE2026"]["D3"] = None
        plan = engine.commit(self.make_batch(book), self.chief)
        self.assertEqual(plan.counts["updated"], 0)
        self.assertEqual(
            IncomingDocument.objects.get(docket_number="2026-06-01-001").initial_remarks,
            "Naomi",
        )

    def test_a_batch_cannot_be_committed_twice(self):
        batch = self.make_batch(incoming_workbook())
        engine.commit(batch, self.chief)
        with self.assertRaises(engine.AlreadyCommitted):
            engine.commit(batch, self.chief)

    def test_imported_records_stay_out_of_the_review_queue(self):
        engine.commit(self.make_batch(incoming_workbook()), self.chief)
        self.assertFalse(IncomingDocument.objects.for_review().exists())
        self.assertFalse(IncomingDocument.objects.open().exists())

        from notifications.service import refresh_standing_notices
        self.assertEqual(refresh_standing_notices()["raised"], 0)


class ViewTests(SyncTestCase):
    def test_upload_preview_commit(self):
        response = self.upload(incoming_workbook())
        batch = SyncBatch.objects.latest("pk")
        self.assertRedirects(response, batch.get_absolute_url())

        page = self.client.get(batch.get_absolute_url())
        self.assertContains(page, "Will be created")
        self.assertContains(page, "No Docket Number")

        self.client.post(reverse("datasync:commit", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, BatchStatus.COMMITTED)
        self.assertEqual(IncomingDocument.objects.count(), 7)
        self.assertContains(self.client.get(reverse("incoming:list")), "2026-01-05-001")

    def test_choosing_sheets_narrows_the_sync(self):
        self.upload(incoming_workbook())
        batch = SyncBatch.objects.latest("pk")
        self.client.post(reverse("datasync:sheets", args=[batch.pk]),
                         {"sheets": ["APRIL"]})
        batch.refresh_from_db()
        self.assertEqual(batch.counts["created"], 2)
        engine.commit(batch, self.chief)
        self.assertEqual(IncomingDocument.objects.count(), 2)

    def test_discard_changes_nothing(self):
        self.upload(incoming_workbook())
        batch = SyncBatch.objects.latest("pk")
        self.client.post(reverse("datasync:discard", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, BatchStatus.DISCARDED)
        self.assertEqual(IncomingDocument.objects.count(), 0)

    def test_encoder_may_not_sync_incoming(self):
        response = self.upload(incoming_workbook(), user=self.encoder)
        self.assertEqual(response.status_code, 200)     # the form refuses the choice
        self.assertFalse(SyncBatch.objects.exists())

    def test_encoder_may_sync_outgoing(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "2026"
        sheet.append(["  ", "DMS Number (Incoming)", "Type of Communication / Documents",
                      "Subject / Title", "Sent to", "Sent Via", "Remarks",
                      "DMS Number (Outgoing)"])
        sheet.append(["LGMED-13 - (DBA) - 2026-01-05-0001", None, "Acknowledgement",
                      "LDRRMF report", "PD Bulabog", "DMS", None, None])
        self.upload(book, module="outgoing", user=self.encoder)
        batch = SyncBatch.objects.latest("pk")
        self.client.post(reverse("datasync:commit", args=[batch.pk]))
        document = OutgoingDocument.objects.get()
        self.assertEqual(document.date_sent, datetime.date(2026, 1, 5))
        self.assertEqual(document.sent_to, "PD Bulabog")
        self.assertContains(self.client.get(reverse("outgoing:list")), "LDRRMF report")

    def test_a_file_that_is_not_a_workbook_is_refused(self):
        self.client.force_login(self.chief)
        bogus = SimpleUploadedFile("register.xlsx", b"not a workbook")
        response = self.client.post(
            reverse("datasync:home"), {"module": "incoming", "file": bogus}
        )
        self.assertContains(response, "could not be read as an Excel workbook")
        self.assertFalse(SyncBatch.objects.exists())


class ClearIncomingTests(SyncTestCase):
    def test_clears_every_incoming_record_without_prompt(self):
        self.upload(incoming_workbook())
        engine.commit(SyncBatch.objects.latest("pk"), self.chief)
        call_command("clear_incoming", "--no-input", stdout=io.StringIO())
        self.assertEqual(IncomingDocument.objects.count(), 0)
        self.assertEqual(IncomingEvent.objects.count(), 0)


# ---------------------------------------------------------------------------
# Outgoing migration
# ---------------------------------------------------------------------------

EN = "–"   # the en dash Word types into " - "
SPEC_HEADERS = ["LGMED CODE", "DNS NUMBER (INCOMING)", "TYPE OF COMMUNICATION/ DOCUMENTS",
                "Subject/ Title", "Sent to", "Sent Via"]


def code(initials, day, number, dash=EN):
    return f"LGMED-13 {dash} ({initials}) {dash} 2026-01-{day:02d}-{number:04d}"


def outgoing_workbook(rows, headers=SPEC_HEADERS):
    book = Workbook()
    sheet = book.active
    sheet.title = "2026"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    return book


class OutgoingMigrationTestCase(SyncTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        from documents.models import DocumentType

        cls.viewer = User.objects.create_user(username="viewer", password="pw",
                                              role=Role.VIEWER)
        cls.document_type = DocumentType.objects.create(name="Memorandum")
        cls.incoming = IncomingDocument.objects.create(
            docket_number="2026-01-02-119", subject="Shear line advisory",
            document_type=cls.document_type, source_office="OCD",
            date_received=datetime.date(2026, 1, 2), created_by=cls.chief,
        )
        # A communication the system already holds, written the system's way.
        cls.stored = OutgoingDocument.objects.create(
            control_code="LGMED-13 - DBA - 2026-01-05-0001",
            subject="Stored subject", sent_to="PD Bulabog",
        )

    def migrate(self, rows, user=None, **options):
        self.upload(outgoing_workbook(rows, **options), module="outgoing", user=user)
        return SyncBatch.objects.latest("pk")

    def statuses(self, batch):
        return {row.row_number: row for row in batch.rows.all()}


class OutgoingMigrationTests(OutgoingMigrationTestCase):
    ROWS = [
        # 2: valid; its DNS number is an incoming document in the system
        [code("RGFJ", 3, 2), "2026-01-02-119", "Memorandum", "Shear line", "PDs/CD", "Email"],
        # 3: the stored record, spelt the register's way - a duplicate
        [code("DBA", 5, 1), None, "Letter", "Changed subject", "Someone else", "DMS"],
        # 4: no LGMED code
        [None, None, "Letter", "A subject", "PD", "DMS"],
        # 5: not a code
        ["END of 2026", None, None, None, None, None],
        # 6: a date that is not a date
        ["LGMED-13 - (KAA) - 2026-02-30-0005", None, "Letter", "Subject", "PD", "DMS"],
        # 7: no subject
        [code("KAA", 6, 6), None, "Letter", None, "PD", "DMS"],
        # 8: valid, but a type nobody has listed
        [code("AFVA", 6, 7), None, "Indorsement", "Q4 CF report", "SILG", "DMS"],
        # 9: "Memo", recognised as Memorandum
        [code("JMR", 7, 8), "DMS No: 2026-01-07-001", "Memo", "Road clearing", "PDs", "Email"],
        # 10: row 2 again, word for word - a duplicate
        [code("RGFJ", 3, 2), "2026-01-02-119", "Memorandum", "Shear line", "PDs/CD", "Email"],
        # 11: row 8's code on a different communication - an error to resolve
        [code("AFVA", 6, 7), None, "Indorsement", "Another report", "SILG", "DMS"],
        # 12: an Excel error value
        [code("DNB", 8, 9), None, "Letter", "#REF!", "PD", "DMS"],
        # 13: blank
        [None, None, None, None, None, None],
    ]

    def test_columns_named_as_in_the_template_are_matched_exactly(self):
        batch = self.migrate(self.ROWS)
        mapping = batch.summary["sheets"][0]["mapping"]
        found = {column["field"]: column["how"] for column in mapping}
        for name in ("control_code", "incoming_reference", "communication_type",
                     "subject", "sent_to", "sent_via"):
            self.assertEqual(found[name], "exact", name)

    def test_a_differently_named_column_is_reported_not_silently_used(self):
        headers = ["Control Code"] + SPEC_HEADERS[1:]
        batch = self.migrate([self.ROWS[0]], headers=headers)
        mapping = {c["field"]: c for c in batch.summary["sheets"][0]["mapping"]}
        self.assertEqual(mapping["control_code"]["how"], "alias")
        self.assertContains(self.client.get(batch.get_absolute_url()),
                            "check this is the same column")

    def test_a_sheet_without_a_required_column_is_not_read(self):
        headers = ["LGMED CODE", "Sent to", "Sent Via", "Remarks"]
        batch = self.migrate([[code("RGFJ", 3, 2), "PD", "DMS", None]], headers=headers)
        sheet = batch.summary["sheets"][0]
        self.assertEqual(sheet["status"], "skipped")
        self.assertIn("Subject / Title", sheet["reason"])

    def test_preview_writes_nothing_and_counts_every_row(self):
        batch = self.migrate(self.ROWS)
        self.assertEqual(OutgoingDocument.objects.count(), 1)
        counts = batch.counts
        self.assertEqual(counts["created"], 3)       # rows 2, 8, 9
        self.assertEqual(counts["duplicates"], 2)    # rows 3, 10
        self.assertEqual(counts["errors"], 6)        # rows 4, 5, 6, 7, 11, 12
        self.assertEqual(counts["total"], 11)
        self.assertEqual(counts["flagged"], 1)       # row 8's type

    def test_each_row_says_what_is_wrong_with_it(self):
        rows = self.statuses(self.migrate(self.ROWS))
        self.assertEqual(rows[2].status, "VALID")
        self.assertEqual(rows[3].status, "DUPLICATE")
        self.assertIn("already in Outgoing Monitoring", rows[3].messages[0])
        self.assertEqual(rows[4].messages, ["No LGMED Code"])
        self.assertIn("not in the LGMED code format", rows[5].messages[0])
        self.assertIn("not a real date", rows[6].messages[0])
        self.assertEqual(rows[7].messages, ["Subject / Title: Missing"])
        self.assertEqual(rows[8].status, "WARNING")
        self.assertIn("Unrecognised type", rows[8].messages[0])
        self.assertEqual(rows[9].values["type"], "Memorandum")
        self.assertEqual(rows[10].status, "DUPLICATE")
        self.assertIn("different communication", rows[11].messages[0])
        self.assertIn("Excel error #REF!", rows[12].messages[0])
        self.assertNotIn(13, rows)                   # blank rows are not records

    def test_import_keeps_codes_exactly_and_skips_everything_else(self):
        batch = self.migrate(self.ROWS)
        self.client.post(reverse("datasync:commit", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, BatchStatus.COMMITTED)

        self.assertEqual(OutgoingDocument.objects.count(), 4)
        imported = OutgoingDocument.objects.get(control_code=code("RGFJ", 3, 2))
        self.assertEqual(imported.control_code, f"LGMED-13 {EN} (RGFJ) {EN} 2026-01-03-0002")
        self.assertEqual(imported.date_sent, datetime.date(2026, 1, 3))
        self.assertEqual(imported.communication_type, "Memorandum")
        self.assertEqual(imported.sent_via, "Email")
        self.assertEqual(imported.created_by, self.chief)
        memo = OutgoingDocument.objects.get(control_code=code("JMR", 7, 8))
        self.assertEqual(memo.communication_type, "Memorandum")
        self.assertEqual(memo.incoming_reference, "2026-01-07-001")
        self.assertEqual(
            OutgoingDocument.objects.get(control_code=code("AFVA", 6, 7)).subject,
            "Q4 CF report",
        )

        # The duplicate left the stored record exactly as it was.
        self.stored.refresh_from_db()
        self.assertEqual(self.stored.subject, "Stored subject")
        self.assertEqual(self.stored.sent_to, "PD Bulabog")

        rows = self.statuses(batch)
        self.assertEqual(rows[2].status, "IMPORTED")
        self.assertEqual(rows[2].record_id, imported.pk)
        self.assertEqual(batch.counts["created"], 3)

    def test_importing_the_same_file_again_creates_nothing(self):
        engine.commit(self.migrate(self.ROWS), self.chief)
        again = self.migrate(self.ROWS)
        self.assertEqual(again.counts["created"], 0)
        self.assertEqual(again.counts["duplicates"], 5)
        engine.commit(again, self.chief)
        self.assertEqual(OutgoingDocument.objects.count(), 4)

    def test_incoming_document_is_referenced_but_not_touched(self):
        before = IncomingDocument.objects.values().get(pk=self.incoming.pk)
        batch = self.migrate(self.ROWS)
        self.assertEqual(self.statuses(batch)[2].values["incoming"]["pk"], self.incoming.pk)
        engine.commit(batch, self.chief)

        self.assertEqual(IncomingDocument.objects.values().get(pk=self.incoming.pk), before)
        record = OutgoingDocument.objects.get(control_code=code("RGFJ", 3, 2))
        self.assertIsNone(record.incoming)
        self.assertEqual(record.incoming_reference, "2026-01-02-119")
        page = self.client.get(record.get_absolute_url())
        self.assertContains(page, self.incoming.get_absolute_url())

    def test_workflow_codes_continue_after_the_migrated_ones(self):
        from incoming import workflow

        year = timezone.localdate().year
        migrated = f"LGMED-13 {EN} (RGFJ) {EN} {year}-01-03-0100"
        engine.commit(
            self.migrate([[migrated, None, "Letter", "Subject", "PD", "DMS"]]), self.chief
        )
        workflow.assign(self.incoming, self.chief, assignee=self.encoder)
        self.incoming.refresh_from_db()
        self.assertTrue(self.incoming.lgmed_code.endswith("-0101"))
        self.assertEqual(self.incoming.status, IncomingStatus.ASSIGNED)

    def test_strict_types_turn_an_unknown_type_into_an_error(self):
        batch = self.migrate(self.ROWS)
        self.client.post(reverse("datasync:sheets", args=[batch.pk]),
                         {"sheets": ["2026"], "strict_types": "1"})
        batch.refresh_from_db()
        self.assertEqual(batch.counts["created"], 2)
        self.assertIn("Unknown document type", self.statuses(batch)[8].messages[0])

    def test_the_migration_is_audited_with_its_counts(self):
        batch = self.migrate(self.ROWS)
        engine.commit(batch, self.chief)
        event = AuditEvent.objects.get(action=Action.SYNC)
        self.assertEqual(event.actor, self.chief)
        self.assertIn("3 imported, 2 duplicates skipped, 6 errors skipped", event.detail)
        self.assertEqual(event.changes["file_name"]["to"], batch.original_name)
        self.assertEqual(event.changes["imported"]["to"], "3")

    def test_a_failure_during_import_saves_nothing(self):
        from unittest import mock

        from django.db import DatabaseError

        batch = self.migrate(self.ROWS)
        with mock.patch("incoming.register.sync_all_outgoing",
                        side_effect=DatabaseError("disk full")):
            response = self.client.post(reverse("datasync:commit", args=[batch.pk]),
                                        follow=True)
        self.assertContains(response, "nothing from this file was saved")
        batch.refresh_from_db()
        self.assertEqual(batch.status, BatchStatus.PREVIEW)
        self.assertEqual(OutgoingDocument.objects.count(), 1)
        self.assertFalse(AuditEvent.objects.filter(action=Action.SYNC).exists())

    def test_error_report_lists_the_rows_left_out_with_reasons(self):
        from openpyxl import load_workbook

        batch = self.migrate(self.ROWS)
        response = self.client.get(reverse("datasync:error_report", args=[batch.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("Import Error Report", response["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][4:10], tuple(SPEC_HEADERS))
        self.assertEqual(len(rows) - 1, 8)          # 6 errors and 2 duplicates
        reasons = {row[1]: row[3] for row in rows[1:]}
        self.assertEqual(reasons[7], "Subject / Title: Missing")
        self.assertIn("#REF!", reasons[12])

    def test_preview_and_result_pages(self):
        batch = self.migrate(self.ROWS)
        page = self.client.get(batch.get_absolute_url())
        self.assertContains(page, "Preview &amp; Validate")
        self.assertContains(page, "Import Valid Records")
        self.assertContains(page, "Records with errors")
        errors = self.client.get(batch.get_absolute_url(), {"status": "error"})
        self.assertContains(errors, "Subject / Title: Missing")
        self.assertNotContains(errors, "Road clearing")

        self.client.post(reverse("datasync:commit", args=[batch.pk]))
        result = self.client.get(batch.get_absolute_url())
        self.assertContains(result, "Migration Completed")
        self.assertContains(result, "Successfully imported")

    def test_viewers_may_not_migrate(self):
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("datasync:home")).status_code, 403)
        batch = self.migrate(self.ROWS)
        self.client.force_login(self.viewer)
        self.assertEqual(
            self.client.get(reverse("datasync:error_report", args=[batch.pk])).status_code,
            403,
        )
        self.client.post(reverse("datasync:commit", args=[batch.pk]))
        self.assertEqual(OutgoingDocument.objects.count(), 1)


class LgmedCodeParsingTests(TestCase):
    def test_register_spellings_of_one_code_are_the_same_code(self):
        from outgoing.codes import canonical

        same = {
            canonical(f"LGMED-13 {EN} (DBA) {EN} 2026-01-05-0001"),
            canonical("LGMED-13 - DBA - 2026-01-05-0001"),
            canonical("LGMED-13-(DBA)-2026-01-05-0001"),
            canonical("lgmed-13 - (dba) - 2026-1-05-1"),
        }
        self.assertEqual(len(same), 1)
        self.assertNotEqual(canonical("LGMED-13 - DBA - 2026-01-05-0001"),
                            canonical("LGMED-13 - DBA - 2026-01-05-0001-A"))

    def test_formats_found_in_the_registers(self):
        from outgoing.codes import parse

        self.assertEqual(parse(f"LGMED-13 {EN} (DR JRRL) {EN} 2023-01-03-0002").initials,
                         "DRJRRL")
        self.assertEqual(parse(f"LGMED-13 {EN} (JTP) {EN} 2022-07-18-1948-A").suffix, "A")
        self.assertTrue(parse(f"LGMED-13 {EN} (DBA)) {EN} 2025-05-28-1259").irregularities)
        self.assertIsNone(parse("LGMED - 13 - (FFJrPB) - 2021-15-2803"))
        self.assertIsNone(parse("NOTHING FOLLOWS! PLEASE ENCODE IN 2022 SHEET"))
