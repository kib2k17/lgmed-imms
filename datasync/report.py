"""
The Import Error Report: every row of a migration that was not imported, and why.

The row's cells are written under the migration template's own headers, so the
report doubles as a worksheet to correct: fix the rows, delete the columns in
front of LGMED CODE if you like, and upload it again. Rows already imported
come back as duplicates and are left alone.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import NOT_IMPORTED


def build(batch):
    """The report as .xlsx bytes."""
    profile = batch.profile
    columns = list(profile.columns)

    book = Workbook()
    sheet = book.active
    sheet.title = "Not imported"
    header = ["Sheet", "Row", "Status", "Reason"] + [
        column.aliases[0] for column in columns
    ]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in batch.rows.filter(status__in=NOT_IMPORTED).order_by("id").iterator():
        values = row.values or {}
        cells = [row.sheet, row.row_number, row.get_status_display(),
                 "\n".join(row.messages or [])]
        for column in columns:
            text = row.key if column.name == profile.key_column else values.get(column.name, "")
            cells.append(text)
        sheet.append(cells)

    widths = [12, 7, 12, 60] + [28] * len(columns)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            # Strings stay strings: a subject that begins with "=" is text, not
            # a formula Excel would run when the report is opened.
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "E2"

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def filename(batch):
    stem = batch.original_name.rsplit(".", 1)[0]
    return f"{stem} - Import Error Report.xlsx"
