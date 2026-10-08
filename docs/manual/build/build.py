"""
Build the LGMED-iMMS System User Manual from its sources.

    python build.py            # Markdown, HTML, PDF, DOCX and PPTX
    python build.py pdf        # only some outputs: md html pdf docx pptx

Sources:  ../src/*.md (chapters, in file-name order)
          ../assets/figures.json and the annotated screenshots (capture.py)
          diagrams.py (workflow diagrams)
Outputs:  ../system-user-manual.{md,html,pdf,docx,pptx}

Tokens understood in the chapters:
    {{figure:<id>}}   the annotated screenshot, numbered caption and legend
    {{diagram:<id>}}  a workflow diagram, numbered like a figure
    {{ref:<id>}}      the number of a figure or diagram
"""

import datetime as dt
import html as htmlmod
import json
import re
import subprocess
import sys
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag

import diagrams

HERE = Path(__file__).resolve().parent
MANUAL = HERE.parent
ROOT = MANUAL.parent.parent
SRC = MANUAL / "src"
ASSETS = MANUAL / "assets"
OUT = MANUAL / "system-user-manual"

TITLE = "SYSTEM USER MANUAL"
SUBTITLE = "Comprehensive Guide to System Modules, Functions, and Procedures"
SYSTEM = "LGMED-iMMS"
SYSTEM_LONG = ("Local Government Monitoring and Evaluation Division - "
               "Information Management and Monitoring System")
OFFICE = "Department of the Interior and Local Government\nRegional Office XIII - Caraga"
MANUAL_VERSION = "1.0"
HISTORY = [
    ("1.0", "8 October 2026", "First edition, generated from the system as built "
     "(commit {commit}). Screenshots captured from the live application on "
     "demonstration data.", "Documentation team / LGMED-iMMS"),
]
TOKEN = re.compile(r"\{\{(figure|diagram|ref):([a-z0-9-]+)\}\}")
CALLOUTS = ("NOTE", "TIP", "IMPORTANT", "WARNING")


def system_version():
    try:
        commit = subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, text=True).strip()
    except Exception:
        commit = "unknown"
    return commit


# --------------------------------------------------------------------- input

def load():
    figures = json.loads((ASSETS / "figures.json").read_text(encoding="utf-8"))
    chapters = [(p.name, p.read_text(encoding="utf-8")) for p in sorted(SRC.glob("*.md"))]
    numbers, order = {}, []
    for _, text in chapters:
        for kind, key in TOKEN.findall(text):
            if kind in ("figure", "diagram") and key not in numbers:
                numbers[key] = len(numbers) + 1
                order.append((kind, key))
    problems = []
    for _, text in chapters:
        for kind, key in TOKEN.findall(text):
            if kind == "figure" and key not in figures:
                problems.append(f"figure {key} has no capture")
            if kind == "diagram" and key not in diagrams.DIAGRAMS:
                problems.append(f"diagram {key} is not defined")
            if kind == "ref" and key not in numbers:
                problems.append(f"ref {key} points at no figure")
    for key, fig in figures.items():
        if key not in numbers:
            problems.append(f"captured figure {key} is not used in the manual")
        if any(not row["found"] for row in fig["legend"]):
            problems.append(f"figure {key} has markers that were not found")
        if not (MANUAL / fig["file"]).exists():
            problems.append(f"figure {key}: image file missing")
    return figures, chapters, numbers, order, problems


def diagram_caption(key):
    return diagrams.DIAGRAMS[key]["title"]


# ------------------------------------------------------------ token expansion

def legend_rows(fig):
    return [row for row in fig["legend"] if row["found"]]


def expand_md(text, figures, numbers):
    def sub(m):
        kind, key = m.groups()
        if kind == "ref":
            return str(numbers[key])
        n = numbers[key]
        if kind == "diagram":
            return (f"![Figure {n}. {diagram_caption(key)}](assets/diagrams/{key}.png)\n\n"
                    f"*Figure {n}. {diagram_caption(key)}*")
        fig = figures[key]
        out = [f"![Figure {n}. {fig['caption']}]({fig['file']})", "",
               f"*Figure {n}. {fig['caption']}.* Signed in as: {fig['role']}."]
        rows = legend_rows(fig)
        if rows:
            out += ["", "| No. | Element | Description |", "|:-:|---|---|"]
            out += [f"| {r['n']} | **{r['label']}** | {r['text']} |" for r in rows]
        return "\n".join(out)
    return TOKEN.sub(sub, text)


def expand_html(text, figures, numbers):
    def sub(m):
        kind, key = m.groups()
        if kind == "ref":
            n = numbers[key]
            return f'<a class="xref" href="#fig-{key}">{n}</a>'
        n = numbers[key]
        if kind == "diagram":
            cap = htmlmod.escape(diagram_caption(key))
            return (f'\n\n<figure class="diagram" id="fig-{key}">'
                    f'<img src="assets/diagrams/{key}.png" alt="{cap}">'
                    f'<figcaption><b>Figure {n}.</b> {cap}</figcaption></figure>\n\n')
        fig = figures[key]
        cap = htmlmod.escape(fig["caption"])
        rows = legend_rows(fig)
        legend = ""
        if rows:
            legend = ('<table class="legend"><thead><tr><th>No.</th><th>Element</th>'
                      '<th>Description</th></tr></thead><tbody>' + "".join(
                          f'<tr><td class="n"><span class="badge">{r["n"]}</span></td>'
                          f'<td><b>{htmlmod.escape(r["label"])}</b></td>'
                          f'<td>{htmlmod.escape(r["text"])}</td></tr>' for r in rows)
                      + "</tbody></table>")
        return (f'\n\n<figure class="shot" id="fig-{key}">'
                f'<img src="{fig["file"]}" alt="{cap}">'
                f'<figcaption><b>Figure {n}.</b> {cap}. '
                f'<span class="role">Signed in as: {htmlmod.escape(fig["role"])}</span>'
                f'</figcaption>{legend}</figure>\n\n')
    return TOKEN.sub(sub, text)


# ------------------------------------------------------------------- markdown

def build_md(figures, chapters, numbers, commit):
    today = dt.date.today().strftime("%d %B %Y").lstrip("0")
    head = [
        f"# {TITLE}", "", f"## {SUBTITLE}", "",
        f"**{SYSTEM}** - {SYSTEM_LONG}  ", OFFICE.replace("\n", "  \n") + "  ",
        f"Manual version {MANUAL_VERSION} - {today} - System build {commit}", "",
        "---", "", "## Document Control and Version History", "",
        "| Version | Date | Description | Prepared by |", "|---|---|---|---|",
    ]
    head += [f"| {v} | {d} | {desc.format(commit=commit)} | {by} |" for v, d, desc, by in HISTORY]
    head += ["", "> This Markdown file is generated by `docs/manual/build/build.py` from "
             "`docs/manual/src/`. Edit the sources, not this file.", "", "---", ""]
    body = "\n\n".join(expand_md(text, figures, numbers) for _, text in chapters)
    (OUT.with_suffix(".md")).write_text("\n".join(head) + "\n" + body, encoding="utf-8")


# ----------------------------------------------------------------------- html

CSS = (HERE / "manual.css").read_text(encoding="utf-8") if (HERE / "manual.css").exists() else ""


def slug(text, used):
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "s"
    base, i = s, 2
    while s in used:
        s, i = f"{base}-{i}", i + 1
    used.add(s)
    return s


def build_html(figures, chapters, numbers, commit, toc_pages=None, embed=False):
    md = markdown.Markdown(extensions=["tables", "sane_lists", "attr_list", "md_in_html"])
    body_html = md.convert("\n\n".join(expand_html(t, figures, numbers) for _, t in chapters))
    soup = BeautifulSoup(body_html, "html.parser")

    if embed:   # the PDF is printed from compressed copies; the HTML keeps the PNGs
        for img in soup.find_all("img"):
            img["src"] = embed_copy(MANUAL / img["src"]).as_uri()

    # Callouts: a blockquote opening with **NOTE:** etc.
    for bq in soup.find_all("blockquote"):
        strong = bq.find("strong")
        if strong and strong.get_text().rstrip(":").strip() in CALLOUTS:
            kind = strong.get_text().rstrip(":").strip().lower()
            bq["class"] = ["callout", kind]
    # "**WARNING:**" inline inside notes paragraphs
    for strong in soup.find_all("strong"):
        if strong.get_text().strip() == "WARNING:" and not strong.find_parent("blockquote"):
            strong["class"] = ["warn-inline"]

    used, toc = set(), []
    for h in soup.find_all(["h1", "h2", "h3"]):
        h["id"] = slug(h.get_text(), used)
        if h.name in ("h1", "h2"):
            toc.append((h.name, h.get_text().strip(), h["id"]))
        if h.name == "h3" and h.get_text().startswith("Procedure"):
            h["class"] = ["procedure"]
    # Procedure label paragraphs
    for p in soup.find_all("p"):
        first = p.contents[0] if p.contents else None
        if isinstance(first, Tag) and first.name == "strong" and first.get_text().rstrip(":") in (
                "Purpose", "Who can do this", "Steps", "Expected Result", "Visual Reference",
                "Important Notes", "Before you start"):
            p["class"] = ["proc-label"]

    today = dt.date.today().strftime("%d %B %Y").lstrip("0")
    logo = (ROOT / "static" / "img" / "dilg-logo.png").as_uri()
    font = (ROOT / "static" / "fonts" / "inter-latin-wght-normal.woff2").as_uri()
    hist = "".join(f"<tr><td>{v}</td><td>{d}</td><td>{htmlmod.escape(desc.format(commit=commit))}</td>"
                   f"<td>{by}</td></tr>" for v, d, desc, by in HISTORY)
    toc_html = "".join(
        f'<li class="toc-{lvl}"><a href="#{hid}"><span class="t">{htmlmod.escape(text)}</span>'
        f'<span class="dots"></span><span class="pg">{(toc_pages or {}).get(hid, "00")}</span></a></li>'
        for lvl, text, hid in toc)
    nfig = len(numbers)
    front = f"""
<section class="cover">
  <img class="logo" src="{logo}" alt="DILG seal">
  <p class="dept">{OFFICE.replace(chr(10), '<br>')}</p>
  <p class="div">Local Government Monitoring and Evaluation Division</p>
  <div class="rule"></div>
  <h1 class="title">{TITLE}</h1>
  <p class="subtitle">{SUBTITLE}</p>
  <p class="system"><b>{SYSTEM}</b><br>{SYSTEM_LONG}</p>
  <table class="meta">
    <tr><th>Manual version</th><td>{MANUAL_VERSION}</td></tr>
    <tr><th>Date</th><td>{today}</td></tr>
    <tr><th>System build</th><td>{commit}</td></tr>
    <tr><th>Classification</th><td>For official use - internal</td></tr>
  </table>
</section>
<section class="front">
  <h2 class="fronth">Document Control and Version History</h2>
  <table class="plain">
    <tr><th>Document title</th><td>{TITLE} - {SUBTITLE}</td></tr>
    <tr><th>System</th><td>{SYSTEM} ({SYSTEM_LONG})</td></tr>
    <tr><th>Office</th><td>{OFFICE.replace(chr(10), ', ')}</td></tr>
    <tr><th>Intended readers</th><td>All LGMED-iMMS users: System Administrators, Administrators (Division Chief), LGMED Staff, Encoders and Viewers</td></tr>
    <tr><th>Basis</th><td>The application as built in the project repository (commit {commit}); every procedure was checked against the source code and the running system</td></tr>
    <tr><th>Figures</th><td>{nfig} annotated screenshots and workflow diagrams</td></tr>
  </table>
  <h3 class="fronth3">Version history</h3>
  <table class="plain"><thead><tr><th>Version</th><th>Date</th><th>Description</th><th>Prepared by</th></tr></thead>
  <tbody>{hist}</tbody></table>
  <h3 class="fronth3">Review and approval</h3>
  <table class="plain sign"><thead><tr><th>Role</th><th>Name</th><th>Signature</th><th>Date</th></tr></thead>
  <tbody><tr><td>Prepared by</td><td></td><td></td><td></td></tr>
  <tr><td>Reviewed by</td><td></td><td></td><td></td></tr>
  <tr><td>Approved by (Division Chief)</td><td></td><td></td><td></td></tr></tbody></table>
  <p class="small">To update this manual, edit the chapter files in <code>docs/manual/src/</code>, recapture screenshots with
  <code>docs/manual/build/capture.py</code>, and regenerate all formats with <code>docs/manual/build/build.py</code>
  (see <code>docs/manual/README.md</code>).</p>
</section>
<section class="toc">
  <h2 class="fronth">Table of Contents</h2>
  <ol class="toc">{toc_html}</ol>
</section>"""
    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>{SYSTEM} {TITLE.title()}</title>
<style>@font-face {{ font-family: Inter; src: url("{font}") format("woff2"); font-weight: 100 900; }}
{CSS}</style></head><body>{front}<main class="content">{soup}</main></body></html>"""
    path = OUT.with_suffix(".html")
    path.write_text(doc, encoding="utf-8")
    return path, toc


# ------------------------------------------------------------------------ pdf

def render_pdf(browser, html_path, pdf_path):
    page = browser.new_page()
    page.goto(html_path.as_uri())
    page.wait_for_load_state("networkidle")
    page.pdf(path=str(pdf_path), format="A4", print_background=True, outline=True, tagged=True,
             margin={"top": "22mm", "bottom": "20mm", "left": "18mm", "right": "18mm"})
    page.close()


def toc_pages_from(pdf_path, toc):
    import pymupdf
    doc = pymupdf.open(pdf_path)
    outline = doc.get_toc(simple=True)
    norm = lambda s: re.sub(r"\s+", " ", re.sub(r"[‒-―-]+", "-", s)).strip().lower()
    found, i = {}, 0
    for lvl, text, hid in toc:
        for j in range(i, len(outline)):
            if norm(outline[j][1]) == norm(text):
                found[hid] = outline[j][2]
                i = j + 1
                break
    # Text search fallback for anything the outline missed: only after the
    # page of the entry before it, so the table of contents itself (which
    # carries every title) is never mistaken for the heading.
    last = 0
    for lvl, text, hid in toc:
        if hid in found:
            last = found[hid]
            continue
        for pno in range(max(last - 1, 4), len(doc)):
            if doc[pno].search_for(text[:40]):
                found[hid] = pno + 1
                break
        last = found.get(hid, last)
    doc.close()
    return found


def decorate_pdf(pdf_path, toc, pages):
    """Running header and footer on every page but the cover; clean bookmarks."""
    import pymupdf
    doc = pymupdf.open(pdf_path)
    total = len(doc)
    for pno in range(1, total):
        page = doc[pno]
        w, h = page.rect.width, page.rect.height
        grey = (0.39, 0.45, 0.55)
        page.insert_text((51, 34), f"{SYSTEM} System User Manual", fontsize=8, color=grey, fontname="helv")
        right = f"Version {MANUAL_VERSION}"
        page.insert_text((w - 51 - pymupdf.get_text_length(right, "helv", 8), 34), right,
                         fontsize=8, color=grey, fontname="helv")
        page.draw_line((51, 40), (w - 51, 40), color=(0.80, 0.84, 0.88), width=0.5)
        page.draw_line((51, h - 40), (w - 51, h - 40), color=(0.80, 0.84, 0.88), width=0.5)
        page.insert_text((51, h - 28), "DILG Regional Office XIII - Caraga | LGMED", fontsize=8,
                         color=grey, fontname="helv")
        label = f"Page {pno + 1} of {total}"
        page.insert_text((w - 51 - pymupdf.get_text_length(label, "helv", 8), h - 28), label,
                          fontsize=8, color=grey, fontname="helv")
    bookmarks = [[1, "Cover", 1], [1, "Document Control and Version History", 2],
                 [1, "Table of Contents", 3]]
    for lvl, text, hid in toc:
        if hid in pages:
            bookmarks.append([1 if lvl == "h1" else 2, text, pages[hid]])
    doc.set_toc(bookmarks)
    doc.set_metadata({"title": f"{SYSTEM} {TITLE.title()}", "subject": SUBTITLE,
                      "author": "DILG Regional Office XIII - Caraga, LGMED",
                      "keywords": "LGMED-iMMS, user manual, DILG Caraga"})
    tmp = pdf_path.with_suffix(".tmp.pdf")
    doc.save(tmp, garbage=3, deflate=True)
    doc.close()
    tmp.replace(pdf_path)


def build_pdf(browser, figures, chapters, numbers, commit):
    pdf = OUT.with_suffix(".pdf")
    html_path, toc = build_html(figures, chapters, numbers, commit, embed=True)
    render_pdf(browser, html_path, pdf)
    pages = toc_pages_from(pdf, toc)
    html_path, toc = build_html(figures, chapters, numbers, commit, toc_pages=pages, embed=True)
    render_pdf(browser, html_path, pdf)
    pages2 = toc_pages_from(pdf, toc)
    if pages2 != pages:   # a TOC entry changed length and moved a page; settle it
        html_path, toc = build_html(figures, chapters, numbers, commit, toc_pages=pages2, embed=True)
        render_pdf(browser, html_path, pdf)
        pages = pages2
    decorate_pdf(pdf, toc, pages)
    build_html(figures, chapters, numbers, commit, toc_pages=pages)   # the viewable HTML
    return pdf, toc, pages


# ------------------------------------------------------- embedded image copies

CACHE = HERE / ".cache" / "embed"


def embed_copy(src):
    """A JPEG copy of an asset for embedding in Word and PowerPoint."""
    from PIL import Image
    src = Path(src)
    out = CACHE / (src.parent.name + "-" + src.stem + ".jpg")
    if not out.exists() or out.stat().st_mtime < src.stat().st_mtime:
        out.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(src) as im:
            im = im.convert("RGB")
            if im.width > 1500:
                im = im.resize((1500, round(im.height * 1500 / im.width)), Image.LANCZOS)
            im.save(out, quality=82, optimize=True)
    return out


# ----------------------------------------------------------------------- docx

def build_docx(commit, toc):
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    html_doc = BeautifulSoup(OUT.with_suffix(".html").read_text(encoding="utf-8"), "html.parser")
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
    sec.different_first_page_header_footer = True
    navy = RGBColor(0x0F, 0x2A, 0x5C)

    st = d.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        st[name].font.name = "Calibri"
        st[name].font.size = Pt(size)
        st[name].font.color.rgb = navy

    def field(paragraph, instr):
        run = paragraph.add_run()
        for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                run._r.append(el)
            elif text == instr:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = instr
                run._r.append(el)
            else:
                t = OxmlElement("w:t")
                t.text = text
                run._r.append(t)

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    # header / footer
    from docx.enum.text import WD_TAB_ALIGNMENT
    usable = sec.page_width - sec.left_margin - sec.right_margin
    hp = sec.header.paragraphs[0]
    from docx.shared import Inches
    for par in (hp, sec.footer.paragraphs[0]):
        stops = par.paragraph_format.tab_stops
        # The Header/Footer styles carry centre and right stops of their own.
        stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
        stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.CLEAR)
        stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    hp.text = f"{SYSTEM} System User Manual\tVersion {MANUAL_VERSION}"
    hp.runs[0].font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    r = fp.add_run("DILG Regional Office XIII - Caraga | LGMED\tPage ")
    r.font.size = Pt(8)
    field(fp, "PAGE")
    fp.add_run(" of ").font.size = Pt(8)
    field(fp, "NUMPAGES")

    # cover
    logo = ROOT / "static" / "img" / "dilg-logo.png"
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(logo), width=Cm(3.5))
    for text, size, bold in ((OFFICE.replace("\n", " - "), 11, False),
                             ("Local Government Monitoring and Evaluation Division", 11, True),
                             ("", 10, False), (TITLE, 30, True), (SUBTITLE, 14, False), ("", 10, False),
                             (SYSTEM, 16, True), (SYSTEM_LONG, 11, False), ("", 10, False),
                             (f"Manual version {MANUAL_VERSION}  |  {dt.date.today():%d %B %Y}  |  System build {commit}", 10, False),
                             ("For official use - internal", 10, False)):
        q = d.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = q.add_run(text)
        run.font.size, run.bold = Pt(size), bold
        if size >= 14:
            run.font.color.rgb = navy
    d.add_page_break()

    # document control
    d.add_heading("Document Control and Version History", 1)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("Version", "Date", "Description", "Prepared by")):
        c.text = h
        shade(c, "E2E8F0")
    for v, dd, desc, by in HISTORY:
        cells = t.add_row().cells
        for c, val in zip(cells, (v, dd, desc.format(commit=commit), by)):
            c.text = val
    d.add_paragraph()
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("Role", "Name", "Signature", "Date")):
        c.text = h
        shade(c, "E2E8F0")
    for role in ("Prepared by", "Reviewed by", "Approved by (Division Chief)"):
        t.add_row().cells[0].text = role
    d.add_page_break()
    d.add_heading("Table of Contents", 1)
    tp = d.add_paragraph()
    field(tp, 'TOC \\o "1-2" \\h \\z \\u')
    note = d.add_paragraph("Right-click the table and choose Update Field to refresh page numbers.")
    note.runs[0].font.size = Pt(8)
    note.runs[0].italic = True
    d.add_page_break()

    # body
    def inline(par, node, bold=False, italic=False):
        for child in node.children:
            if isinstance(child, NavigableString):
                txt = re.sub(r"\s+", " ", str(child))
                if txt.strip() or txt == " ":
                    run = par.add_run(txt)
                    run.bold, run.italic = bold or None, italic or None
            elif child.name in ("strong", "b"):
                inline(par, child, True, italic)
            elif child.name in ("em", "i"):
                inline(par, child, bold, True)
            elif child.name == "code":
                run = par.add_run(child.get_text())
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
            elif child.name == "br":
                par.add_run().add_break()
            elif child.name in ("span", "a"):
                inline(par, child, bold, italic)
            else:
                inline(par, child, bold, italic)

    def add_list(node, ordered, level=0):
        # Numbers are written into the text rather than left to Word's list
        # numbering, which would run on from one procedure to the next.
        for i, li in enumerate(node.find_all("li", recursive=False), 1):
            if ordered:
                par = d.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.9 + 0.8 * level)
                par.paragraph_format.first_line_indent = Cm(-0.6)
                par.add_run(f"{i}.	")
                par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9 + 0.8 * level))
            else:
                par = d.add_paragraph(style="List Bullet 2" if level else "List Bullet")
            par.paragraph_format.space_after = Pt(2)
            for child in li.children:
                if isinstance(child, Tag) and child.name in ("ul", "ol"):
                    continue
                if isinstance(child, Tag) and child.name == "p":
                    inline(par, child)
                elif isinstance(child, NavigableString):
                    if str(child).strip():
                        par.add_run(re.sub(r"\s+", " ", str(child)))
                else:
                    inline(par, BeautifulSoup(str(child), "html.parser"))
            for sub in li.find_all(["ul", "ol"], recursive=False):
                add_list(sub, sub.name == "ol", level + 1)

    def add_table(node, widths=None, legend=False):
        rows = node.find_all("tr")
        ncols = max(len(r.find_all(["td", "th"])) for r in rows)
        t = d.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r in rows:
            cells = t.add_row().cells
            for c, src in zip(cells, r.find_all(["td", "th"])):
                c.paragraphs[0].text = ""
                inline(c.paragraphs[0], src, bold=src.name == "th")
                for run in c.paragraphs[0].runs:
                    run.font.size = Pt(9)
                if src.name == "th":
                    shade(c, "E2E8F0")
                if legend and src.get("class") == ["n"]:
                    shade(c, "EA580C")
                    for run in c.paragraphs[0].runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            for row in t.rows:
                for c, wdt in zip(row.cells, widths):
                    c.width = wdt
        d.add_paragraph()

    from PIL import Image
    main = html_doc.find("main")
    first_h1 = True
    for node in main.children:
        if not isinstance(node, Tag):
            continue
        if node.name == "h1":
            if not first_h1:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_h1 = False
            d.add_heading(node.get_text(), 1)
        elif node.name in ("h2", "h3", "h4"):
            d.add_heading(node.get_text(), int(node.name[1]))
        elif node.name == "p":
            par = d.add_paragraph()
            inline(par, node)
        elif node.name in ("ul", "ol"):
            add_list(node, node.name == "ol")
        elif node.name == "table":
            add_table(node)
        elif node.name == "blockquote":
            t = d.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            kind = (node.get("class") or ["", "note"])[-1]
            shade(t.rows[0].cells[0], {"note": "EFF6FF", "tip": "F0FDF4", "important": "FFF7ED",
                                       "warning": "FEF2F2"}.get(kind, "F8FAFC"))
            cell = t.rows[0].cells[0]
            cell.paragraphs[0].text = ""
            first = True
            for para in node.find_all(["p", "li"]):
                par = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                inline(par, para)
            d.add_paragraph()
        elif node.name == "figure":
            img = node.find("img")
            src = MANUAL / img["src"]   # PNG path in the saved HTML
            with Image.open(src) as im:
                wpx, hpx = im.size
            width = Cm(16.5)
            if hpx / wpx > 1.25:          # tall figures: limit by height
                width = Cm(min(16.5, 22.0 * wpx / hpx))
            par = d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(embed_copy(src)), width=width)
            cap = d.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(cap, node.find("figcaption"))
            for run in cap.runs:
                run.font.size = Pt(9)
                run.italic = True
            legend = node.find("table")
            if legend:
                add_table(legend, widths=(Cm(1.3), Cm(4.2), Cm(11)), legend=True)
        elif node.name == "hr":
            continue
    # Ask Word to refresh the table of contents and page fields on opening.
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    d.settings.element.append(upd)
    d.core_properties.title = f"{SYSTEM} {TITLE.title()}"
    d.core_properties.subject = SUBTITLE
    d.core_properties.author = "DILG Regional Office XIII - Caraga, LGMED"
    d.save(OUT.with_suffix(".docx"))


# ----------------------------------------------------------------------- pptx

def build_pptx(figures, chapters, numbers, commit):
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    navy, orange, slate = RGBColor(0x0F, 0x2A, 0x5C), RGBColor(0xEA, 0x58, 0x0C), RGBColor(0x33, 0x41, 0x55)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, value, size=14, bold=False, colour=slate, align=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        lines = value if isinstance(value, list) else [value]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = colour
            run.font.name = "Calibri"
            if align:
                p.alignment = align
        return box

    def band(slide, colour=navy, h=0.9):
        shp = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = colour
        shp.line.fill.background()
        return shp

    def footer(slide, n, dark=False):
        c = RGBColor(0xCB, 0xD5, 0xE1) if dark else RGBColor(0x64, 0x74, 0x8B)
        text(slide, 0.4, 7.05, 9, 0.3, f"{SYSTEM} System User Manual | DILG Regional Office XIII - Caraga | LGMED",
             9, colour=c)
        text(slide, 11.9, 7.05, 1.1, 0.3, str(n), 9, colour=c)

    # cover
    s = prs.slides.add_slide(blank)
    band(s, navy, 7.5)
    s.shapes.add_picture(str(ROOT / "static/img/dilg-logo.png"), Inches(0.8), Inches(0.8), height=Inches(1.4))
    text(s, 0.8, 2.6, 11.5, 0.9, TITLE, 44, True, RGBColor(255, 255, 255))
    text(s, 0.8, 3.5, 11.5, 0.6, SUBTITLE, 22, False, RGBColor(0xCB, 0xD5, 0xE1))
    text(s, 0.8, 4.5, 11.5, 0.9, [SYSTEM, SYSTEM_LONG], 16, False, RGBColor(255, 255, 255))
    text(s, 0.8, 6.2, 11.5, 0.6, f"{OFFICE.replace(chr(10), ' - ')}  |  Version {MANUAL_VERSION}  |  "
         f"{dt.date.today():%d %B %Y}  |  Build {commit}", 12, False, RGBColor(0xCB, 0xD5, 0xE1))
    count = 1

    # contents
    chapter_titles = []
    for _, txt in chapters:
        chapter_titles += re.findall(r"^# (.+)$", txt, re.M)
    s = prs.slides.add_slide(blank)
    band(s)
    text(s, 0.5, 0.18, 12, 0.6, "Contents", 26, True, RGBColor(255, 255, 255))
    half = (len(chapter_titles) + 1) // 2
    text(s, 0.6, 1.2, 6, 5.6, chapter_titles[:half], 15)
    text(s, 6.9, 1.2, 6, 5.6, chapter_titles[half:], 15)
    count += 1
    footer(s, count)

    for name, txt in chapters:
        # walk tokens and headings in order
        for m in re.finditer(r"^(#{1,2}) (.+)$|\{\{(figure|diagram):([a-z0-9-]+)\}\}", txt, re.M):
            if m.group(1):
                level, title = len(m.group(1)), m.group(2)
                if level == 1 or (name.startswith("10-") and re.match(r"10\.\d+ ", title)):
                    s = prs.slides.add_slide(blank)
                    band(s, navy if level == 1 else RGBColor(0x1E, 0x3A, 0x8A), 7.5)
                    text(s, 0.8, 3.0, 11.7, 1.2, title, 36 if level == 1 else 32, True, RGBColor(255, 255, 255))
                    count += 1
                    footer(s, count, dark=True)
                continue
            kind, key = m.group(3), m.group(4)
            n = numbers[key]
            if kind == "diagram":
                src, caption, legend = ASSETS / "diagrams" / f"{key}.png", diagrams.DIAGRAMS[key]["title"], []
            else:
                fig = figures[key]
                src, caption, legend = MANUAL / fig["file"], fig["caption"], legend_rows(fig)
            s = prs.slides.add_slide(blank)
            band(s)
            text(s, 0.4, 0.15, 12.5, 0.65, f"Figure {n}. {caption}", 22, True, RGBColor(255, 255, 255))
            with Image.open(src) as im:
                wpx, hpx = im.size
            area_w = 8.4 if legend else 12.3
            area_h = 5.9
            scale = min(area_w / wpx, area_h / hpx)
            w, h = wpx * scale, hpx * scale
            s.shapes.add_picture(str(embed_copy(src)), Inches(0.5 + (area_w - w) / 2), Inches(1.05 + (area_h - h) / 2),
                                 Inches(w), Inches(h))
            if legend:
                y = 1.1
                size = 12 if len(legend) <= 6 else 10.5 if len(legend) <= 8 else 9.5
                step = (5.8 / max(len(legend), 1))
                for row in legend:
                    dot = s.shapes.add_shape(9, Inches(9.15), Inches(y + 0.02), Inches(0.36), Inches(0.36))
                    dot.fill.solid()
                    dot.fill.fore_color.rgb = orange
                    dot.line.fill.background()
                    dot.text_frame.text = str(row["n"])
                    para = dot.text_frame.paragraphs[0]
                    para.runs[0].font.size, para.runs[0].font.bold = Pt(11), True
                    para.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    text(s, 9.6, y - 0.05, 3.5, step, [row["label"], row["text"]], size)
                    box = s.shapes[-1].text_frame.paragraphs[0].runs[0]
                    box.font.bold = True
                    box.font.color.rgb = navy
                    y += step
            if kind == "figure":
                text(s, 0.5, 6.75, 8.4, 0.3, f"Signed in as: {figures[key]['role']}", 10,
                     colour=RGBColor(0x64, 0x74, 0x8B))
            count += 1
            footer(s, count)

    prs.core_properties.title = f"{SYSTEM} {TITLE.title()}"
    prs.save(OUT.with_suffix(".pptx"))


# ----------------------------------------------------------------------- main

def main(targets):
    from playwright.sync_api import sync_playwright

    commit = system_version()
    figures, chapters, numbers, order, problems = load()
    if problems:
        print("Problems found:\n  " + "\n  ".join(problems))
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome")
        diagrams.render_all(browser, ASSETS / "diagrams")
        if "md" in targets:
            build_md(figures, chapters, numbers, commit)
            print("written", OUT.with_suffix(".md").name)
        if "pdf" in targets or "html" in targets or "docx" in targets:
            pdf, toc, pages = build_pdf(browser, figures, chapters, numbers, commit)
            print("written", OUT.with_suffix(".html").name, "and", pdf.name)
        browser.close()
    if "docx" in targets:
        build_docx(commit, toc)
        print("written", OUT.with_suffix(".docx").name)
    if "pptx" in targets:
        build_pptx(figures, chapters, numbers, commit)
        print("written", OUT.with_suffix(".pptx").name)
    print(f"{len(numbers)} figures and diagrams; {len(problems)} problem(s).")


if __name__ == "__main__":
    main(set(sys.argv[1:]) or {"md", "html", "pdf", "docx", "pptx"})
